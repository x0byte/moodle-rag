from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .base import Page


def _shape_text(shape) -> list[str]:
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return [text for child in shape.shapes for text in _shape_text(child)]
    if getattr(shape, "has_table", False) and shape.has_table:
        return [
            " | ".join(cell.text.strip() for cell in row.cells) for row in shape.table.rows
        ]
    if shape.has_text_frame:
        return [shape.text_frame.text.strip()]
    return []


def extract_pptx(path: Path) -> list[Page]:
    """One page per slide: its text, tables and speaker notes."""
    pages = []
    for number, slide in enumerate(Presentation(str(path)).slides, start=1):
        parts = [text for shape in slide.shapes for text in _shape_text(shape) if text]
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None and notes.text.strip():
                parts.append(f"Speaker notes: {notes.text.strip()}")
        pages.append(Page(number=number, text="\n".join(parts)))
    return pages
